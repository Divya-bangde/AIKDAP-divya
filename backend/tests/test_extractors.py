"""Extractor tests (Sprint 12.1).

Fixtures are generated with the same libraries the extractors read
with a different library on the writing side wherever one is
available (python-docx/openpyxl/python-pptx are read/write round-trip
libraries; PDF authoring uses reportlab since pypdf has no text-layout
API of its own — see requirements.txt for why). Every fixture is
built in-memory, deterministic, and small; nothing here touches the
database or a running service, so these tests exercise `extractors.py`
in isolation from the rest of the pipeline.
"""

import io

import pytest

from app.modules.assets.processing.extractors import (
    CsvExtractor,
    DocxExtractor,
    ExtractedDocument,
    ExtractionFailedError,
    ExtractionNotSupportedError,
    HtmlExtractor,
    ImageExtractor,
    JsonExtractor,
    OcrRequiredError,
    PdfExtractor,
    PlainTextExtractor,
    PptxExtractor,
    UnsupportedExtractor,
    XlsxExtractor,
    get_text_extractor,
)

pytestmark = pytest.mark.asyncio


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _make_pdf(pages: list[str]) -> bytes:
    """A real, correctly-structured multi-page PDF with the given text
    per page. An empty string produces a genuinely blank page (no text
    objects at all), for testing mixed/all-empty-page scenarios."""
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(300, 300))
    for text in pages:
        if text:
            pdf.drawString(20, 150, text)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _make_text_image(text: str | None, *, size: tuple[int, int] = (900, 200)) -> bytes:
    """A PNG containing real, readable rendered text (via a proper TTF
    font, not PIL's low-quality default bitmap font) -- or, if `text`
    is `None`, a plain solid-color image with no text at all, for
    testing "OCR ran, found nothing" scenarios with a genuine image
    (as opposed to a page with no image at all)."""
    from PIL import Image as PILImage
    from PIL import ImageDraw, ImageFont

    img = PILImage.new("RGB", size, color="white")
    if text:
        font = ImageFont.truetype(
            "/usr/local/lib/python3.12/site-packages/reportlab/fonts/Vera.ttf", 28
        )
        ImageDraw.Draw(img).text((20, 20), text, fill="black", font=font)
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def _make_scanned_pdf(pages: list[str | None]) -> bytes:
    """A PDF where each page is a full-page IMAGE (a real scanned-PDF
    shape), not a blank page with no content at all. `None` produces a
    page whose image has no readable text (OCR should find nothing);
    a string produces a page whose image contains that real, readable
    text, rendered via a proper font."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(900, 200))
    for text in pages:
        image_bytes = _make_text_image(text)
        pdf.drawImage(ImageReader(io.BytesIO(image_bytes)), 0, 0, width=900, height=200)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _make_rotated_text_image(text: str, angle: int, *, size: tuple[int, int] = (900, 200)) -> bytes:
    """Like `_make_text_image`, but rotated by `angle` degrees after
    rendering -- simulates a scanner feeding a page in sideways/upside
    down (Sprint 12.5 Phase 3)."""
    from PIL import Image as PILImage

    png_bytes = _make_text_image(text, size=size)
    image = PILImage.open(io.BytesIO(png_bytes))
    rotated = image.rotate(angle, expand=True, fillcolor="white")
    buffer = io.BytesIO()
    rotated.save(buffer, format="PNG")
    return buffer.getvalue()


def _make_scanned_pdf_from_images(pages: list[bytes]) -> bytes:
    """Like `_make_scanned_pdf`, but takes pre-built image bytes per
    page directly, so a caller can supply a rotated image."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(900, 200))
    for image_bytes in pages:
        pdf.drawImage(ImageReader(io.BytesIO(image_bytes)), 0, 0, width=900, height=200)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _make_encrypted_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret", owner_password="secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_docx(paragraphs: list[tuple[str, str | None]], table_rows: list[list[str]] | None = None) -> bytes:
    """`paragraphs` is (text, style_name_or_None); style "Heading 1" etc.
    marks a heading. `table_rows[0]` is the header row."""
    import docx

    document = docx.Document()
    for text, style in paragraphs:
        document.add_paragraph(text, style=style)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx(sheets: dict[str, list[list[str]]]) -> bytes:
    """`sheets["Name"][0]` is the header row."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_pptx(slides: list[tuple[str | None, list[str]]]) -> bytes:
    """`slides` is a list of (title_or_None, body_text_lines)."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title + content
    for title, body_lines in slides:
        slide = presentation.slides.add_slide(layout)
        if title:
            slide.shapes.title.text = title
        else:
            slide.shapes.title.text = ""
        if body_lines and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.text = body_lines[0]
            for line in body_lines[1:]:
                paragraph = body.add_paragraph()
                paragraph.text = line
        elif body_lines:
            textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(2))
            textbox.text_frame.text = "\n".join(body_lines)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Registry / dispatch
# ---------------------------------------------------------------------------


def test_registry_dispatches_every_supported_mime_type():
    expectations = {
        "text/plain": PlainTextExtractor,
        "text/markdown": PlainTextExtractor,
        "application/json": JsonExtractor,
        "text/csv": CsvExtractor,
        "application/pdf": PdfExtractor,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocxExtractor,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": XlsxExtractor,
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": PptxExtractor,
        "text/html": HtmlExtractor,
    }
    for mime_type, expected_class in expectations.items():
        assert isinstance(get_text_extractor(mime_type), expected_class)


def test_registry_falls_back_to_unsupported_for_legacy_office_formats():
    for mime_type in ("application/msword", "application/vnd.ms-excel", "application/vnd.ms-powerpoint"):
        assert isinstance(get_text_extractor(mime_type), UnsupportedExtractor)


async def test_unsupported_extractor_raises_not_swallows():
    extractor = get_text_extractor("application/octet-stream")
    with pytest.raises(ExtractionNotSupportedError):
        await extractor.extract(b"anything")


# ---------------------------------------------------------------------------
# Plain text / Markdown
# ---------------------------------------------------------------------------


async def test_plain_text_extraction():
    doc = await PlainTextExtractor().extract("Hello\n\nWorld".encode())
    assert doc.full_text == "Hello\n\nWorld"
    assert doc.has_extractable_text


async def test_plain_text_collapses_excessive_blank_lines():
    doc = await PlainTextExtractor().extract(b"Para one\n\n\n\n\nPara two")
    assert doc.full_text == "Para one\n\nPara two"


async def test_plain_text_empty_content_yields_no_units():
    doc = await PlainTextExtractor().extract(b"   \n\n  ")
    assert doc.units == []
    assert not doc.has_extractable_text


async def test_plain_text_rejects_binary_content_mislabeled_as_text():
    """Sprint 12.4: a file whose bytes are mostly non-text (a binary
    file mislabeled `text/plain`, bypassing or predating the
    extension/MIME cross-check) must fail extraction rather than
    silently index PDF-syntax-as-if-it-were-prose. Deterministic binary
    content (not a real PDF) so the test doesn't depend on a specific
    fixture library's compression ratio."""
    binary_content = bytes(i % 256 for i in range(4000))
    with pytest.raises(ExtractionFailedError, match="does not appear to be valid text"):
        await PlainTextExtractor().extract(binary_content)


async def test_plain_text_accepts_genuine_text_with_occasional_accents():
    """The garbage-ratio gate must not false-positive on ordinary text
    containing a modest amount of non-ASCII (accented names, curly
    quotes, etc.) -- only content that is *mostly* non-text."""
    text = ("Café résumé naïve — a perfectly normal sentence with a few accents. " * 20).encode(
        "utf-8"
    )
    doc = await PlainTextExtractor().extract(text)
    assert doc.has_extractable_text
    assert "Café" in doc.full_text


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------


async def test_json_extraction_pretty_prints_and_preserves_structure():
    content = b'{"name":"Acme","revenue":50000,"regions":["West","East"]}'
    doc = await JsonExtractor().extract(content)
    assert '"name": "Acme"' in doc.full_text
    assert '"revenue": 50000' in doc.full_text
    assert "West" in doc.full_text


async def test_json_invalid_content_fails_not_silently_passed_through():
    with pytest.raises(ExtractionFailedError):
        await JsonExtractor().extract(b"{not valid json")


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


async def test_csv_rows_are_header_labelled_not_raw_commas():
    content = b"Product,Region,Revenue\nLaptop,West,50000\nPhone,East,30000\n"
    doc = await CsvExtractor().extract(content)
    assert len(doc.units) == 2
    assert doc.units[0].text == "Row: 1\nProduct: Laptop\nRegion: West\nRevenue: 50000"
    assert doc.units[1].text == "Row: 2\nProduct: Phone\nRegion: East\nRevenue: 30000"


async def test_csv_header_only_yields_no_units():
    doc = await CsvExtractor().extract(b"Product,Region,Revenue\n")
    assert doc.units == []


# ---------------------------------------------------------------------------
# PDF — the sprint's priority #1
# ---------------------------------------------------------------------------


async def test_pdf_single_page_extraction():
    pdf_bytes = _make_pdf(["Hello from page one"])
    doc = await PdfExtractor().extract(pdf_bytes)
    assert len(doc.units) == 1
    assert doc.units[0].page_number == 1
    assert "Hello from page one" in doc.units[0].text


async def test_pdf_multi_page_extraction_preserves_page_numbers():
    pdf_bytes = _make_pdf(["First page content", "Second page content", "Third page content"])
    doc = await PdfExtractor().extract(pdf_bytes)
    assert [unit.page_number for unit in doc.units] == [1, 2, 3]
    assert "First page content" in doc.units[0].text
    assert "Second page content" in doc.units[1].text
    assert "Third page content" in doc.units[2].text


async def test_pdf_middle_and_final_page_content_individually_addressable():
    pdf_bytes = _make_pdf([f"Page {i} unique marker {i}" for i in range(1, 6)])
    doc = await PdfExtractor().extract(pdf_bytes)
    by_page = {unit.page_number: unit.text for unit in doc.units}
    assert "marker 3" in by_page[3]
    assert "marker 5" in by_page[5]


async def test_pdf_mixed_empty_pages_recorded_as_warning_not_failure():
    pdf_bytes = _make_pdf(["Real content", "", "More real content"])
    doc = await PdfExtractor().extract(pdf_bytes)
    # Only the two non-empty pages become units; the blank one is
    # skipped, not fabricated, and noted in warnings.
    assert [unit.page_number for unit in doc.units] == [1, 3]
    assert any("1 of 3 pages" in warning for warning in doc.warnings)


async def test_pdf_all_pages_empty_and_no_images_fails_cleanly():
    """Sprint 12.5: a page with no text AND no embedded image at all
    (a genuinely blank/vector-only page, as these `_make_pdf` fixtures
    with an empty string produce -- no `drawImage` call, nothing in
    the content stream) has nothing for OCR to work with. With OCR
    available on this deployment, that is now reported as a genuine
    extraction failure, not `OCR_REQUIRED` -- `OCR_REQUIRED` means "a
    scan exists but this deployment can't read it", which is a
    different, narrower claim than "this page has no content of any
    kind". See `test_ocr_*` below for the real scanned-page case
    (an image with unreadable content) using `_make_scanned_pdf`."""
    pdf_bytes = _make_pdf(["", "", ""])
    with pytest.raises(ExtractionFailedError, match="OCR was attempted"):
        await PdfExtractor().extract(pdf_bytes)


async def test_pdf_majority_empty_no_images_fails_despite_some_real_text():
    """Sprint 12.3's original regression guard, updated for Sprint
    12.5: a document that is mostly blank (no text, no images) but has
    a genuine text page must not be silently indexed on that small
    fraction alone. Blank-with-no-image pages can't be OCR'd, so this
    is a `FAILED` extraction now that OCR exists and was genuinely
    available -- see the OCR section below for the real scanned-page
    equivalent (majority pages ARE images, just unreadable ones)."""
    pdf_bytes = _make_pdf(["Cover page with real text", "", "", "", ""])
    with pytest.raises(ExtractionFailedError, match="have no usable text even after OCR"):
        await PdfExtractor().extract(pdf_bytes)


async def test_pdf_minority_empty_pages_still_only_a_warning():
    """The majority-empty threshold must not regress the existing
    'mostly fine, a few blank pages' case: below the threshold, this
    stays a warning and the readable pages are still indexed."""
    pdf_bytes = _make_pdf(["Real content one", "Real content two", "Real content three", ""])
    doc = await PdfExtractor().extract(pdf_bytes)
    assert len(doc.units) == 3
    assert any("1 of 4 pages" in warning for warning in doc.warnings)


async def test_pdf_corrupted_bytes_fail_not_silently_pass():
    with pytest.raises(ExtractionFailedError):
        await PdfExtractor().extract(b"%PDF-1.4\nthis is not a real pdf structure")


async def test_pdf_encrypted_reports_password_protected():
    pdf_bytes = _make_encrypted_pdf()
    with pytest.raises(ExtractionFailedError, match="password"):
        await PdfExtractor().extract(pdf_bytes)


async def test_pdf_empty_bytes_fail_cleanly():
    with pytest.raises(ExtractionFailedError):
        await PdfExtractor().extract(b"")


# ---------------------------------------------------------------------------
# OCR (Sprint 12.5) -- real embedded images, not blank/no-image pages.
# `_make_scanned_pdf` embeds a genuine PNG (via `drawImage`) on every
# page, matching a real scanned PDF's shape; `_make_pdf` above never
# does, which is exactly why the two fixture builders produce
# different outcomes for a page with no readable text.
# ---------------------------------------------------------------------------


async def test_ocr_recovers_readable_text_from_a_scanned_page():
    pdf_bytes = _make_scanned_pdf(["Invoice number 4471, total due $250.00."])
    doc = await PdfExtractor().extract(pdf_bytes)
    assert len(doc.units) == 1
    assert doc.units[0].page_number == 1
    assert doc.units[0].extraction_method == "ocr"
    assert "Invoice number 4471" in doc.units[0].text


async def test_ocr_multi_page_scanned_document_preserves_page_numbers():
    pdf_bytes = _make_scanned_pdf(
        ["First scanned page content here.", "Second scanned page content here."]
    )
    doc = await PdfExtractor().extract(pdf_bytes)
    assert [unit.page_number for unit in doc.units] == [1, 2]
    assert all(unit.extraction_method == "ocr" for unit in doc.units)
    assert "First scanned page" in doc.units[0].text
    assert "Second scanned page" in doc.units[1].text


async def test_ocr_mixed_native_and_scanned_pages_both_recovered():
    """The Phase 3 diagram's exact case: page 1 native text, page 2
    scanned. Both must appear, correctly attributed."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(900, 200))
    pdf.drawString(20, 100, "Native text page one content.")
    pdf.showPage()
    image_bytes = _make_text_image("Scanned page two content.")
    pdf.drawImage(ImageReader(io.BytesIO(image_bytes)), 0, 0, width=900, height=200)
    pdf.showPage()
    pdf.save()
    pdf_bytes = buffer.getvalue()

    doc = await PdfExtractor().extract(pdf_bytes)
    assert len(doc.units) == 2
    by_page = {unit.page_number: unit for unit in doc.units}
    assert by_page[1].extraction_method == "native"
    assert "Native text page one" in by_page[1].text
    assert by_page[2].extraction_method == "ocr"
    assert "Scanned page two" in by_page[2].text


async def test_ocr_unreadable_scanned_page_fails_not_silently_empty():
    """A page that IS an image (unlike the blank/no-image fixtures
    above) but whose image has no recognizable text -- OCR runs,
    finds nothing, and the document fails honestly rather than
    completing with zero content."""
    pdf_bytes = _make_scanned_pdf([None, None])
    with pytest.raises(ExtractionFailedError, match="OCR was attempted"):
        await PdfExtractor().extract(pdf_bytes)


async def test_ocr_disabled_reports_ocr_required_not_failed(monkeypatch):
    """When OCR is unavailable on this deployment (disabled by
    configuration here; the same path a missing `tesseract` binary
    would take), a real scanned page must still be reported as
    `OCR_REQUIRED` -- the pre-12.5 behavior, unchanged for a
    deployment that genuinely cannot OCR."""
    from app.core.config import settings

    monkeypatch.setattr(settings, "ocr_enabled", False)
    pdf_bytes = _make_scanned_pdf(["Some readable scanned text."])
    with pytest.raises(OcrRequiredError):
        await PdfExtractor().extract(pdf_bytes)


async def test_ocr_page_limit_is_respected(monkeypatch):
    """The `ocr_max_pages_per_document` cap must actually bound how
    many pages get OCR'd -- verified by setting it to 2 against a
    3-scanned-page document (so the 1 capped-out page is a minority,
    staying under the majority-empty threshold and proving the cap
    itself works rather than triggering the separate failure path)."""
    from app.core.config import settings

    monkeypatch.setattr(settings, "ocr_max_pages_per_document", 2)
    pdf_bytes = _make_scanned_pdf(
        ["Page one scanned text.", "Page two scanned text.", "Page three scanned text."]
    )
    doc = await PdfExtractor().extract(pdf_bytes)
    assert len(doc.units) == 2
    assert any("OCR page limit" in warning for warning in doc.warnings)


# ---------------------------------------------------------------------------
# OSD-based rotation correction (Sprint 12.5 Phase 3)
#
# Phase 2's audit measured that a 90-/180-degree-rotated scanned page
# previously produced completely wrong text that still passed the
# confidence gate (40.5, just above the 40.0 threshold) -- a silent
# correctness bug, not a rejection. These tests exercise the fix
# through the real `PdfExtractor` page loop, not a mocked OCR call.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("angle", [90, 180, 270])
async def test_ocr_rotated_scanned_page_recovers_correct_text(angle):
    text = "Invoice number 4471, total due $250.00."
    image_bytes = _make_rotated_text_image(text, angle)
    pdf_bytes = _make_scanned_pdf_from_images([image_bytes])

    doc = await PdfExtractor().extract(pdf_bytes)

    assert len(doc.units) == 1
    assert doc.units[0].page_number == 1
    assert doc.units[0].extraction_method == "ocr"
    assert text in doc.units[0].text, (
        f"rotation={angle}: expected correct text, got {doc.units[0].text!r} "
        "-- OSD correction did not recover the page"
    )


async def test_ocr_normal_orientation_unaffected_by_osd_correction():
    """0 degrees (no rotation needed) must behave exactly as before --
    OSD detecting 'rotate: 0' must be a no-op, not a regression."""
    text = "Invoice number 4471, total due $250.00."
    pdf_bytes = _make_scanned_pdf([text])

    doc = await PdfExtractor().extract(pdf_bytes)

    assert len(doc.units) == 1
    assert text in doc.units[0].text
    assert doc.units[0].extraction_method == "ocr"


async def test_ocr_rotation_correction_applies_only_to_pages_that_need_ocr():
    """Mixed PDF: page 1 native text, page 2 a 90-degree-rotated scan.
    OSD/rotation correction must never touch the native page -- it only
    ever runs inside the OCR fallback path.

    Uses the same longer, realistic fixture text as the other rotation
    tests -- Tesseract OSD needs a reasonable amount of text to
    classify orientation confidently at all (measured: a short 4-word
    image raises its own "too few characters" error, the same
    inconclusive-fallback path `test_ocr_osd_failure_falls_back_to_
    unrotated_ocr` covers deliberately); this test is about provenance
    (correction only touching the OCR'd page), not about that
    known, separate OSD limitation.
    """
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    text = "Invoice number 4471, total due $250.00."
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(900, 200))
    pdf.drawString(20, 100, "Native text page one content.")
    pdf.showPage()
    rotated_image = _make_rotated_text_image(text, 90)
    pdf.drawImage(ImageReader(io.BytesIO(rotated_image)), 0, 0, width=900, height=200)
    pdf.showPage()
    pdf.save()
    pdf_bytes = buffer.getvalue()

    doc = await PdfExtractor().extract(pdf_bytes)

    assert len(doc.units) == 2
    by_page = {unit.page_number: unit for unit in doc.units}
    assert by_page[1].extraction_method == "native"
    assert "Native text page one" in by_page[1].text
    assert by_page[2].extraction_method == "ocr"
    assert text in by_page[2].text


async def test_ocr_osd_genuinely_inconclusive_on_very_short_text_fails_honestly():
    """A real (not mocked) OSD limitation, discovered while writing
    these tests: Tesseract OSD raises its own 'too few characters'
    error on very short/sparse rotated text and cannot correct it, so
    OCR proceeds on the still-rotated image. That, in turn, produces
    text that fails the EXISTING word-confidence gate (`ocr.py`'s
    `ocr_min_confidence` check, unchanged by this sprint) -- so the
    page is correctly reported as unreadable rather than silently
    indexed as wrong content. This is materially better than the
    pre-fix behavior (garbled text that quietly passed the gate): a
    clean, honest failure instead of silent corruption. Confirms the
    OSD fallback does not crash even when it cannot help."""
    rotated_image = _make_rotated_text_image("Hi there.", 90)
    pdf_bytes = _make_scanned_pdf_from_images([rotated_image])

    with pytest.raises(ExtractionFailedError, match="OCR was attempted"):
        await PdfExtractor().extract(pdf_bytes)


async def test_ocr_osd_failure_falls_back_to_unrotated_ocr(monkeypatch):
    """If OSD itself raises (simulating a Tesseract OSD failure, e.g.
    unreadable/too-few-characters), OCR must still succeed on the
    original image rather than failing the whole page."""
    import pytesseract

    def _raising_osd(*args, **kwargs):
        raise pytesseract.TesseractError(1, "Too few characters. Skipping this page.")

    monkeypatch.setattr(pytesseract, "image_to_osd", _raising_osd)

    text = "Invoice number 4471, total due $250.00."
    pdf_bytes = _make_scanned_pdf([text])

    doc = await PdfExtractor().extract(pdf_bytes)

    assert len(doc.units) == 1
    assert text in doc.units[0].text


async def test_ocr_osd_low_confidence_suggestion_is_ignored(monkeypatch):
    """A structurally 'successful' but low-confidence OSD read (the
    real random-noise measurement was orientation_conf=0.13) must not
    be trusted -- the image must be OCR'd unrotated."""
    import pytesseract

    def _low_confidence_osd(*args, **kwargs):
        return {"rotate": 90, "orientation_conf": 0.13, "script": "Latin", "script_conf": 0.3}

    monkeypatch.setattr(pytesseract, "image_to_osd", _low_confidence_osd)

    text = "Invoice number 4471, total due $250.00."
    pdf_bytes = _make_scanned_pdf([text])  # normal, unrotated scan

    doc = await PdfExtractor().extract(pdf_bytes)

    # If the low-confidence "rotate: 90" suggestion had been trusted,
    # this correctly-oriented page would have been rotated INTO being
    # wrong and OCR would fail to recover the text.
    assert len(doc.units) == 1
    assert text in doc.units[0].text


# ---------------------------------------------------------------------------
# Standalone image extraction (Sprint 12.5)
# ---------------------------------------------------------------------------


async def test_image_extractor_recovers_readable_text():
    png_bytes = _make_text_image("AIKDAP was created in 2026.")
    doc = await ImageExtractor().extract(png_bytes)
    assert len(doc.units) == 1
    assert doc.units[0].page_number is None
    assert doc.units[0].extraction_method == "ocr"
    assert "AIKDAP was created in 2026" in doc.units[0].text


async def test_image_extractor_unreadable_image_fails_cleanly():
    png_bytes = _make_text_image(None)
    with pytest.raises(ExtractionFailedError):
        await ImageExtractor().extract(png_bytes)


async def test_image_extractor_corrupt_bytes_fail_cleanly():
    with pytest.raises(ExtractionFailedError):
        await ImageExtractor().extract(b"not a real image file")


async def test_image_extractor_ocr_disabled_reports_ocr_required(monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "ocr_enabled", False)
    png_bytes = _make_text_image("Some readable text.")
    with pytest.raises(OcrRequiredError):
        await ImageExtractor().extract(png_bytes)


async def test_image_extractor_recovers_readable_text_from_jpeg():
    """JPEG regression (Sprint 12.5 Phase 3): unchanged behavior,
    exercised explicitly rather than only via registry dispatch."""
    from PIL import Image as PILImage

    png_bytes = _make_text_image("AIKDAP was created in 2026.")
    image = PILImage.open(io.BytesIO(png_bytes)).convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=90)

    doc = await ImageExtractor().extract(buffer.getvalue())
    assert len(doc.units) == 1
    assert doc.units[0].page_number is None
    assert "AIKDAP was created in 2026" in doc.units[0].text


async def test_image_registry_dispatches_png_and_jpeg():
    assert isinstance(get_text_extractor("image/png"), ImageExtractor)
    assert isinstance(get_text_extractor("image/jpeg"), ImageExtractor)


# ---------------------------------------------------------------------------
# GIF / WEBP (Sprint 12.5 Phase 3)
#
# Phase 2 measured that both formats decode via Pillow and OCR
# correctly when called directly, but had no registry entry -- always
# falling through to `UnsupportedExtractor`. These tests exercise the
# real registry dispatch, not a direct `ocr_image()` call.
# ---------------------------------------------------------------------------


def _make_text_gif(text: str, *, animated: bool = False) -> bytes:
    from PIL import Image as PILImage

    png_bytes = _make_text_image(text)
    frame = PILImage.open(io.BytesIO(png_bytes)).convert("RGB")
    buffer = io.BytesIO()
    if animated:
        second_frame = _make_text_image("A completely different second frame.")
        second = PILImage.open(io.BytesIO(second_frame)).convert("RGB")
        frame.save(buffer, format="GIF", save_all=True, append_images=[second], duration=200, loop=0)
    else:
        frame.save(buffer, format="GIF")
    return buffer.getvalue()


def _make_text_webp(text: str) -> bytes:
    from PIL import Image as PILImage

    png_bytes = _make_text_image(text)
    image = PILImage.open(io.BytesIO(png_bytes)).convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="WEBP")
    return buffer.getvalue()


async def test_gif_registry_dispatches_to_image_extractor():
    assert isinstance(get_text_extractor("image/gif"), ImageExtractor)


async def test_webp_registry_dispatches_to_image_extractor():
    assert isinstance(get_text_extractor("image/webp"), ImageExtractor)


async def test_static_gif_recovers_readable_text():
    gif_bytes = _make_text_gif("AIKDAP was created in 2026.")
    doc = await get_text_extractor("image/gif").extract(gif_bytes)
    assert len(doc.units) == 1
    assert doc.units[0].page_number is None
    assert doc.units[0].extraction_method == "ocr"
    assert "AIKDAP was created in 2026" in doc.units[0].text


async def test_webp_recovers_readable_text():
    webp_bytes = _make_text_webp("AIKDAP was created in 2026.")
    doc = await get_text_extractor("image/webp").extract(webp_bytes)
    assert len(doc.units) == 1
    assert doc.units[0].page_number is None
    assert doc.units[0].extraction_method == "ocr"
    assert "AIKDAP was created in 2026" in doc.units[0].text


async def test_animated_gif_uses_first_frame_only():
    """Documented, tested behavior (Sprint 12.5 Phase 3): an animated
    GIF is OCR'd on its first frame only -- never treated as a
    multi-page document, and never silently mixing text from later
    frames into the result."""
    gif_bytes = _make_text_gif("AIKDAP was created in 2026.", animated=True)
    doc = await get_text_extractor("image/gif").extract(gif_bytes)
    assert len(doc.units) == 1
    assert "AIKDAP was created in 2026" in doc.units[0].text
    assert "completely different second frame" not in doc.units[0].text


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


async def test_docx_paragraphs_and_headings_extracted_with_section():
    docx_bytes = _make_docx(
        [
            ("Introduction", "Heading 1"),
            ("This is the intro paragraph.", None),
            ("Methodology", "Heading 1"),
            ("This describes the method.", None),
        ]
    )
    doc = await DocxExtractor().extract(docx_bytes)
    by_text = {unit.text: unit.section for unit in doc.units}
    assert by_text["This is the intro paragraph."] == "Introduction"
    assert by_text["This describes the method."] == "Methodology"


async def test_docx_table_rendered_header_labelled():
    docx_bytes = _make_docx(
        [("Sales Data", "Heading 1")],
        table_rows=[["Product", "Revenue"], ["Laptop", "50000"], ["Phone", "30000"]],
    )
    doc = await DocxExtractor().extract(docx_bytes)
    table_units = [u.text for u in doc.units if u.text.startswith("Table")]
    assert "Table 1, row 1: Product=Laptop | Revenue=50000" in table_units
    assert "Table 1, row 2: Product=Phone | Revenue=30000" in table_units


async def test_docx_corrupted_fails_cleanly():
    with pytest.raises(ExtractionFailedError):
        await DocxExtractor().extract(b"not a real docx file")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


async def test_xlsx_rows_labelled_and_grouped_by_sheet():
    xlsx_bytes = _make_xlsx(
        {
            "Sales": [["Product", "Region", "Revenue"], ["Laptop", "West", 50000]],
            "Inventory": [["Item", "Count"], ["Widget", 12]],
        }
    )
    doc = await XlsxExtractor().extract(xlsx_bytes)
    by_sheet = {unit.sheet_name: unit.text for unit in doc.units}
    assert (
        by_sheet["Sales"]
        == "Sheet: Sales\nRow: 1\nProduct: Laptop\nRegion: West\nRevenue: 50000"
    )
    assert by_sheet["Inventory"] == "Sheet: Inventory\nRow: 1\nItem: Widget\nCount: 12"


async def test_xlsx_corrupted_fails_cleanly():
    with pytest.raises(ExtractionFailedError):
        await XlsxExtractor().extract(b"not a real xlsx file")


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


async def test_pptx_slides_extracted_with_title_and_page_number():
    pptx_bytes = _make_pptx(
        [
            ("Welcome", ["Opening remarks"]),
            ("Q3 Results", ["Revenue grew 12%", "Costs held flat"]),
        ]
    )
    doc = await PptxExtractor().extract(pptx_bytes)
    assert [unit.page_number for unit in doc.units] == [1, 2]
    assert doc.units[0].section == "Welcome"
    assert "Opening remarks" in doc.units[0].text
    assert doc.units[1].section == "Q3 Results"
    assert "Revenue grew 12%" in doc.units[1].text


async def test_pptx_corrupted_fails_cleanly():
    with pytest.raises(ExtractionFailedError):
        await PptxExtractor().extract(b"not a real pptx file")


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------


async def test_html_rejects_binary_content_mislabeled_as_html():
    """Sprint 12.4: same garbage-ratio gate as PlainTextExtractor,
    applied to HtmlExtractor since it decodes with the same
    errors="replace" leniency."""
    binary_content = bytes(i % 256 for i in range(4000))
    with pytest.raises(ExtractionFailedError, match="does not appear to be valid HTML"):
        await HtmlExtractor().extract(binary_content)


async def test_html_sections_by_heading_scripts_and_styles_stripped():
    html = b"""
    <html><head><style>body{color:red}</style></head>
    <body>
      <h1>Overview</h1>
      <p>This is the overview paragraph.</p>
      <h2>Details</h2>
      <p>This is the details paragraph.</p>
      <script>alert('should not appear')</script>
    </body></html>
    """
    doc = await HtmlExtractor().extract(html)
    full = doc.full_text
    assert "should not appear" not in full
    assert "color:red" not in full
    by_section = {unit.section: unit.text for unit in doc.units}
    assert by_section["Overview"] == "This is the overview paragraph."
    assert by_section["Details"] == "This is the details paragraph."


async def test_html_with_no_structural_tags_falls_back_to_whole_text():
    doc = await HtmlExtractor().extract(b"<div>Just some text in a div</div>")
    assert "Just some text in a div" in doc.full_text


# ---------------------------------------------------------------------------
# ExtractedDocument helper properties
# ---------------------------------------------------------------------------


def test_extracted_document_full_text_joins_units_with_blank_line():
    from app.modules.assets.processing.extractors import ExtractedUnit

    doc = ExtractedDocument(units=[ExtractedUnit(text="A"), ExtractedUnit(text="B")])
    assert doc.full_text == "A\n\nB"


def test_extracted_document_empty_has_no_extractable_text():
    assert not ExtractedDocument(units=[]).has_extractable_text
