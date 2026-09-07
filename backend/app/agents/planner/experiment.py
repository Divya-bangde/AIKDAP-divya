"""Experiment Plan Engine (Sprint 16 Phase 6).

Turns a research equation, a `ResearchDocumentUnderstanding`, and/or
user-supplied test cases into a structured, reviewable EXPERIMENT PLAN.

Execution boundary (Part P, absolute rule #10/#12): nothing in this
module runs generated code, trains a model, or produces an observed
result. It only classifies, validates, and organizes -- the researcher
reviews and edits the plan before any future execution phase exists.

Deterministic-first (Part V/W): equation parsing and constraint
derivation are SymPy, numeric work is plain Python, and constraint
enforcement never defers to the LLM. The LLM (existing `LLMGateway`,
Gemini A -> Gemini B -> Groq -> OpenRouter, no second router) is used
only for the one thing deterministic code cannot do: interpreting a
variable's likely research role from prose when the paper/context
evidence hierarchy (Part D) runs out -- and even then it is the last,
lowest-confidence rung, never a substitute for real evidence.
"""

from __future__ import annotations

import csv
import io
import itertools
import json
import re
import signal
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone

import sympy
from openpyxl import load_workbook
from sympy.parsing.sympy_parser import (
    implicit_multiplication_application,
    parse_expr,
    standard_transformations,
)

from app.modules.research.experiment_schemas import (
    ConstraintSource,
    ExperimentConstraint,
    ExperimentInput,
    ExperimentOutput,
    ExperimentOutputValue,
    ExperimentTestCase,
    ExperimentVariable,
    ExperimentVariant,
    MutabilityEvidenceSource,
    MutabilityStatus,
    OutputKind,
    RunPlanStatus,
    VariableRole,
)
from app.modules.research.schemas import ResearchCertainty, ResearchVariable

_SYMPY_TRANSFORMATIONS = standard_transformations + (implicit_multiplication_application,)

# SECURITY (Part U/Y.9): `sympy.parse_expr`/`sympify` parses by handing
# the (transformed) token stream to Python's real `eval()`. Two
# independent layers are required, not one:
#
# 1. `_SAFE_GLOBAL_DICT` blocks `__builtins__` in the eval namespace --
#    without this, `__import__('os').system(...)` in an equation or
#    constraint string genuinely executes (verified empirically while
#    building this module: it printed real shell output).
# 2. `_assert_expression_is_safe` rejects the input BEFORE parsing if it
#    contains anything outside a plain-math character set, a dunder
#    ('__'), or an attribute-access dot -- because blocking
#    `__builtins__` alone is NOT sufficient. The classic sandbox-escape
#    gadget `().__class__.__bases__[0].__subclasses__()` needs no
#    builtin at all (it walks live object attributes already reachable
#    from a literal) and was confirmed, empirically, to still return
#    the full list of loaded Python classes -- including
#    `subprocess.Popen` -- through `global_dict={'__builtins__': {}}`
#    alone. The character/pattern allowlist below is what actually
#    closes that path; the restricted global_dict is defense in depth
#    for anything the allowlist doesn't anticipate (e.g. a bare
#    `eval(...)`/`open(...)` call, which contains no forbidden
#    character but must still resolve to nothing).
_SAFE_GLOBAL_DICT: dict = {}
exec("from sympy import *", _SAFE_GLOBAL_DICT)  # noqa: S102 - trusted, hardcoded string, not user input
_SAFE_GLOBAL_DICT["__builtins__"] = {}

# `\s` in a non-ASCII regex matches Unicode whitespace (e.g. U+00A0
# NBSP, U+2028 LINE SEPARATOR), not just space/tab -- confirmed during
# this module's Phase 6 security review to be broader than intended,
# though harmless in practice (Python's own tokenizer rejects those
# characters downstream, verified empirically). Restricting to ASCII
# space/tab removes the gap outright rather than relying on downstream
# behavior to keep catching it.
_SAFE_EXPRESSION_CHARSET = re.compile(r"^[A-Za-z0-9_ \t.+\-*/()<>=!,]+$")
_DUNDER_PATTERN = re.compile(r"__")
_ATTRIBUTE_ACCESS_DOT = re.compile(r"(?<![0-9])\.(?![0-9])")

# SECURITY (Phase 6 security review, denial-of-service class): a
# character-level allowlist stops object-model escapes but does
# nothing about *mathematically legitimate* input that is
# computationally explosive -- confirmed empirically that
# `parse_equation("Y = 9**9**9**9")` hangs past 20 seconds (9 raised to
# a tower whose value has more digits than fit in memory), and that
# `factorial(factorial(20))` hangs the same way even through the
# already-`evaluate=False` constraint path, because SymPy's automatic
# evaluation of concrete-integer function calls happens independently
# of the `evaluate` flag on `parse_expr`. Because this deployment runs
# uvicorn as a single process with no worker pool, one such request
# freezes the process for every user, not just the caller. A length
# cap closes the unbounded-input half of this (also fixes an
# unhandled `RecursionError` a 200k-character expression produced,
# verified empirically); the wall-clock guard below closes the
# short-but-explosive half, since no static character or function-name
# blocklist can enumerate every SymPy function capable of eager,
# expensive evaluation.
_MAX_EXPRESSION_LENGTH = 500
_PARSE_TIMEOUT_SECONDS = 5


class EquationParseError(ValueError):
    """The equation could not be parsed. Never silently swallowed --
    the caller must surface this to the user rather than guess."""


class _ExpressionTimeoutError(EquationParseError):
    """Raised when parsing/evaluating an expression exceeds
    `_PARSE_TIMEOUT_SECONDS` -- treated as a parse failure, not a
    crash, so it surfaces to the caller as an ordinary 400 rather than
    an unhandled exception."""


@contextmanager
def _bounded_parse_time():
    """Wall-clock budget around a single SymPy parse/evaluate call.
    `signal.alarm` is Unix/main-thread only (this module only ever
    runs inside the Linux backend/worker containers, never on
    Windows) -- if `SIGALRM` is unavailable for any reason, this
    degrades to a no-op rather than raising, so the length cap and
    character allowlist remain as defense-in-depth instead of the
    module becoming uncallable.
    """
    alarm = getattr(signal, "alarm", None)
    sigalrm = getattr(signal, "SIGALRM", None)
    if alarm is None or sigalrm is None:
        yield
        return

    def _on_timeout(signum, frame):
        raise _ExpressionTimeoutError(
            f"Expression took longer than {_PARSE_TIMEOUT_SECONDS}s to parse/evaluate -- "
            "likely a pathological expression (e.g. a large exponent tower or nested "
            "factorial). Rejected rather than left to run."
        )

    previous_handler = signal.signal(sigalrm, _on_timeout)
    alarm(_PARSE_TIMEOUT_SECONDS)
    try:
        yield
    finally:
        alarm(0)
        signal.signal(sigalrm, previous_handler)


def _assert_expression_is_safe(expression: str) -> None:
    """Reject anything that is not plainly arithmetic/comparison
    syntax, before it ever reaches SymPy's `eval`-based parser. Blocks
    dunder attribute access (`__class__`, `__subclasses__`, ...),
    quoting, brackets, semicolons, any character outside a minimal
    math-expression charset, and any expression over
    `_MAX_EXPRESSION_LENGTH` -- not merely a `__import__`/`eval`
    blocklist, which is exactly the approach that failed against the
    `__subclasses__` gadget chain during this module's own security
    testing.
    """
    if not expression or not expression.strip():
        raise EquationParseError("Expression is empty.")
    if len(expression) > _MAX_EXPRESSION_LENGTH:
        raise EquationParseError(
            f"Expression is too long ({len(expression)} characters; the limit is "
            f"{_MAX_EXPRESSION_LENGTH})."
        )
    if not _SAFE_EXPRESSION_CHARSET.match(expression):
        raise EquationParseError(
            "Expression contains characters that are not valid in a mathematical "
            "expression or constraint."
        )
    if _DUNDER_PATTERN.search(expression):
        raise EquationParseError("Expression must not contain '__'.")
    if _ATTRIBUTE_ACCESS_DOT.search(expression):
        raise EquationParseError(
            "Expression must not contain attribute-access syntax ('.') outside of a "
            "decimal number."
        )


class EquationParseResult:
    """Deterministic decomposition of one equation. No mutability
    decision is made here -- that is a separate, evidence-driven step
    (Part D), never inferred from the math alone."""

    def __init__(
        self,
        raw_expression: str,
        output_name: str | None,
        rhs_expr: sympy.Expr,
        input_symbols: list[str],
        parameter_symbols: list[str],
        constraints: list[ExperimentConstraint],
    ) -> None:
        self.raw_expression = raw_expression
        self.output_name = output_name
        self.rhs_expr = rhs_expr
        self.input_symbols = input_symbols
        self.parameter_symbols = parameter_symbols
        self.constraints = constraints


def parse_equation(expression: str, *, source_reference=None) -> EquationParseResult:
    """Parse `Y = (a*X + b) / (c*X + d)` (or `(aX+b)/(cX+d)` without an
    explicit `Y =`) into an output name, free symbols, and mathematically
    derived constraints (Part E/F).

    A symbol is provisionally classified INPUT if it also appears as a
    bare single-letter symbol conventionally used for the independent
    variable (heuristic: the LAST free symbol in alphabetical position
    among single-letter symbols is not a safe signal by itself -- so
    this function does NOT guess input vs. parameter; it returns every
    RHS free symbol as a `parameter_symbols` candidate and leaves INPUT
    classification to the caller, who has the paper's `ResearchVariable`
    roles to draw on). Raises `EquationParseError` on invalid syntax --
    never returns a partially-parsed result.
    """
    raw = expression.strip()
    if not raw:
        raise EquationParseError("Equation expression is empty.")

    output_name: str | None = None
    rhs_text = raw
    if "=" in raw:
        lhs, _, rhs_text = raw.partition("=")
        lhs = lhs.strip()
        if lhs and re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", lhs):
            output_name = lhs
        rhs_text = rhs_text.strip()

    _assert_expression_is_safe(rhs_text)
    try:
        with _bounded_parse_time():
            rhs_expr = parse_expr(
                rhs_text,
                global_dict=_SAFE_GLOBAL_DICT,
                transformations=_SYMPY_TRANSFORMATIONS,
                evaluate=False,
            )
    except (SyntaxError, TypeError, sympy.SympifyError, RecursionError) as exc:
        raise EquationParseError(f"Could not parse equation: {exc}") from exc

    free_symbols = sorted((str(s) for s in rhs_expr.free_symbols))
    if output_name in free_symbols:
        # e.g. "Y = Y + 1" -- not a well-formed forward equation.
        raise EquationParseError(
            f"Output symbol '{output_name}' also appears on the right-hand side."
        )

    constraints: list[ExperimentConstraint] = []
    try:
        with _bounded_parse_time():
            numerator, denominator = sympy.fraction(sympy.together(rhs_expr))
            if denominator != 1 and not denominator.is_number:
                constraints.append(
                    ExperimentConstraint(
                        expression=f"({sympy.sstr(denominator)}) != 0",
                        description="Derived from the equation's denominator: division by "
                        "zero must be avoided.",
                        source=ConstraintSource.MATHEMATICAL_CONSTRAINT,
                        source_reference=source_reference,
                    )
                )
    except (TypeError, ValueError, RecursionError):
        # Not every valid SymPy expression is a clean fraction (e.g. no
        # division at all) -- absence of a denominator constraint is not
        # an error, just nothing to report.
        pass
    except _ExpressionTimeoutError:
        # Denominator/constraint derivation is a convenience, not the
        # equation's own validity -- a pathological RHS that already
        # parsed above but is too expensive to reduce to a single
        # fraction should not block plan creation, just skip deriving
        # this one constraint.
        pass

    return EquationParseResult(
        raw_expression=raw,
        output_name=output_name,
        rhs_expr=rhs_expr,
        input_symbols=[],
        parameter_symbols=free_symbols,
        constraints=constraints,
    )


_ROLE_KEYWORDS: dict[VariableRole, tuple[str, ...]] = {
    VariableRole.INPUT: ("input", "independent variable", "feature", "predictor"),
    VariableRole.OUTPUT: ("output", "dependent variable", "target", "response", "label"),
    VariableRole.PARAMETER: ("parameter", "hyperparameter", "coefficient", "setting"),
    VariableRole.CONSTANT: ("constant", "fixed value"),
    VariableRole.DERIVED_QUANTITY: ("derived", "computed", "calculated from"),
}


def classify_variable_role(research_variable: ResearchVariable) -> VariableRole:
    """Map a `ResearchVariable.role` free-text field onto the closed
    `VariableRole` vocabulary via keyword matching (Part C). Returns
    UNKNOWN rather than guessing when no keyword matches -- an unmapped
    role must not silently become VARIABLE by default."""
    role_text = (research_variable.role or "").lower()
    for role, keywords in _ROLE_KEYWORDS.items():
        if any(keyword in role_text for keyword in keywords):
            return role
    return VariableRole.UNKNOWN


def analyze_mutability(
    *,
    name: str,
    paper_statement: str | None = None,
    experiment_context: str | None = None,
    workspace_mutable_hint: bool | None = None,
    research_understanding_role: VariableRole | None = None,
    llm_suggested_mutable: bool | None = None,
) -> tuple[MutabilityStatus, str, MutabilityEvidenceSource]:
    """Walk the Part D evidence hierarchy in order, stopping at the
    first source that actually says something. Never returns MUTABLE
    or FIXED without a source_reference-carrying reason attached by the
    caller -- this function returns the reason text, the caller is
    responsible for keeping it paired with the variable.
    """
    if paper_statement:
        lowered = paper_statement.lower()
        if any(w in lowered for w in ("fixed", "held constant", "not varied", "kept constant")):
            return (
                MutabilityStatus.FIXED,
                f"The paper explicitly states this is fixed: \"{paper_statement}\"",
                MutabilityEvidenceSource.PAPER_STATEMENT,
            )
        if any(w in lowered for w in ("varied", "tuned", "swept", "explored", "hyperparameter")):
            return (
                MutabilityStatus.MUTABLE,
                f"The paper explicitly treats this as varied/tuned: \"{paper_statement}\"",
                MutabilityEvidenceSource.PAPER_STATEMENT,
            )

    if experiment_context:
        lowered = experiment_context.lower()
        if "test case" in lowered or "ablation" in lowered or "sweep" in lowered:
            return (
                MutabilityStatus.MUTABLE,
                f"Named in the experiment/test-case context as varied: \"{experiment_context}\"",
                MutabilityEvidenceSource.EXPERIMENT_CONTEXT,
            )

    if workspace_mutable_hint is not None:
        return (
            MutabilityStatus.MUTABLE if workspace_mutable_hint else MutabilityStatus.FIXED,
            "Set explicitly by the current workspace/application state.",
            MutabilityEvidenceSource.WORKSPACE_STATE,
        )

    if research_understanding_role is not None:
        if research_understanding_role == VariableRole.PARAMETER:
            return (
                MutabilityStatus.MUTABLE,
                "Explicitly identified as a training/experiment parameter in the structured "
                "research understanding.",
                MutabilityEvidenceSource.RESEARCH_UNDERSTANDING,
            )
        if research_understanding_role in (VariableRole.CONSTANT, VariableRole.DERIVED_QUANTITY):
            return (
                MutabilityStatus.FIXED,
                f"Classified as {research_understanding_role.value} in the structured research "
                "understanding; not treated as a tunable quantity unless the research "
                "explicitly says otherwise.",
                MutabilityEvidenceSource.RESEARCH_UNDERSTANDING,
            )

    if llm_suggested_mutable is not None:
        return (
            MutabilityStatus.MUTABLE if llm_suggested_mutable else MutabilityStatus.FIXED,
            "No explicit paper/context/workspace evidence was found; this is an LLM "
            "interpretation only and should be treated as low-confidence.",
            MutabilityEvidenceSource.LLM_INTERPRETATION,
        )

    return (
        MutabilityStatus.UNKNOWN,
        "No evidence (paper statement, experiment context, workspace state, or structured "
        "understanding) establishes whether this can safely be changed.",
        MutabilityEvidenceSource.NONE,
    )


def build_variables_from_equation(
    parse_result: EquationParseResult,
    *,
    research_variables: dict[str, ResearchVariable] | None = None,
    known_inputs: list[str] | None = None,
    source_reference=None,
) -> tuple[list[ExperimentVariable], ExperimentOutput | None, list[ExperimentInput]]:
    """Turn a parsed equation into `ExperimentVariable`s (Part E).

    Every RHS symbol is `role=UNKNOWN, mutable=UNKNOWN` unless a
    matching `ResearchVariable` (by name) supplies real evidence, or the
    caller explicitly names it in `known_inputs` -- the equation's math
    alone can identify the OUTPUT (the `Y =` left-hand side) but cannot
    tell an input apart from a parameter among the right-hand-side
    symbols (in `Y = (aX+b)/(cX+d)`, nothing about the algebra says X is
    "the input" rather than a). Guessing that would violate Part Y.16
    ("do not invent variable mutability/role"), so it is left UNKNOWN
    unless the paper, the caller, or the user says so.
    """
    research_variables = research_variables or {}
    known_input_set = set(known_inputs or [])
    variables: list[ExperimentVariable] = []
    inputs: list[ExperimentInput] = []

    for symbol_name in parse_result.parameter_symbols:
        matched = research_variables.get(symbol_name)
        if symbol_name in known_input_set:
            role = VariableRole.INPUT
        elif matched:
            role = classify_variable_role(matched)
        else:
            role = VariableRole.UNKNOWN
        understanding_role = role if (matched or symbol_name in known_input_set) else None
        mutable, reason, evidence = analyze_mutability(
            name=symbol_name,
            paper_statement=matched.description if matched else None,
            research_understanding_role=understanding_role,
        )
        if role == VariableRole.INPUT:
            inputs.append(
                ExperimentInput(
                    name=symbol_name, value="", type="number", source_reference=source_reference
                )
            )
            continue
        variables.append(
            ExperimentVariable(
                name=symbol_name,
                role=role,
                mutable=mutable,
                mutability_reason=reason,
                mutability_evidence=evidence,
                unit=matched.unit if matched else None,
                certainty=matched.confidence if matched else ResearchCertainty.UNKNOWN,
                source_reference=source_reference,
            )
        )

    output = None
    if parse_result.output_name:
        output = ExperimentOutput(
            name=parse_result.output_name,
            type="number",
            description=f"Output of: {parse_result.raw_expression}",
            source_reference=source_reference,
        )

    return variables, output, inputs


def check_all_constraints(
    variables: list[ExperimentVariable], constraints: list[ExperimentConstraint]
) -> list[str]:
    """Validate every constraint whose referenced symbols all have a
    numeric `current_value` among `variables`. A constraint that
    mentions a variable not yet given a value is skipped, not failed --
    it cannot be evaluated, which is not the same as being violated.
    Returns one human-readable message per violated constraint; an
    empty list means every checkable constraint passed. This is the
    deterministic gate Part W requires before any future execution
    phase -- e.g. it is what rejects `learning_rate = -0.1` against
    `learning_rate > 0`.
    """
    numeric_values: dict[str, float] = {}
    for var in variables:
        if var.current_value is None:
            continue
        try:
            numeric_values[var.name] = float(var.current_value)
        except ValueError:
            continue

    # Same local_dict requirement as `validate_numeric_constraint`: a
    # known multi-letter variable name (e.g. "threshold") must be
    # declared up front or the implicit-multiplication transform will
    # misread it as a product of single-letter symbols. Built from
    # every variable on the plan, not just the ones with a value yet --
    # a constraint mentioning an as-yet-unvalued variable must still
    # parse correctly so it can be correctly SKIPPED, not silently
    # misparsed into unrelated single-letter symbols.
    local_dict = {var.name: sympy.Symbol(var.name) for var in variables}

    violations: list[str] = []
    for constraint in constraints:
        try:
            _assert_expression_is_safe(constraint.expression)
            with _bounded_parse_time():
                expr = parse_expr(
                    constraint.expression,
                    local_dict=local_dict,
                    global_dict=_SAFE_GLOBAL_DICT,
                    transformations=_SYMPY_TRANSFORMATIONS,
                    evaluate=False,
                )
        except (SyntaxError, TypeError, sympy.SympifyError, RecursionError, EquationParseError):
            # An unsafe or unparseable constraint is inert here, not
            # violated: it is never evaluated, and it never blocks a
            # plan on its own -- the same treatment as any other
            # constraint this function cannot check (Part W).
            continue
        needed = {str(s) for s in expr.free_symbols}
        if not needed.issubset(numeric_values.keys()):
            continue
        try:
            satisfied = validate_numeric_constraint(
                constraint.expression, {k: numeric_values[k] for k in needed}
            )
        except EquationParseError:
            continue
        if not satisfied:
            violations.append(
                f"Constraint violated: {constraint.expression} ({constraint.description})"
            )
    return violations


def validate_numeric_constraint(expression: str, values: dict[str, float]) -> bool:
    """Deterministically check one constraint expression (e.g.
    'c*X + d != 0' or 'learning_rate > 0') against concrete values.
    Returns True if satisfied. Never uses the LLM (Part W) -- this is
    the last line of defense before any future execution phase, and
    correctness here cannot depend on a model's judgment.
    """
    # A `local_dict` naming every known variable up front is required
    # here: without it, `implicit_multiplication_application` (needed
    # so equation shorthand like "aX" parses as a*X) will just as
    # happily read a genuine multi-letter variable name it doesn't
    # recognize -- e.g. "threshold" -- as an implicit product of single
    # -letter symbols (d*e*h*h*l*o*r*s*t). Declaring the real symbols
    # in advance is what tells the parser "threshold" is one name, not
    # nine multiplied together.
    _assert_expression_is_safe(expression)
    local_dict = {name: sympy.Symbol(name) for name in values}
    # The whole parse-substitute-evaluate sequence shares one timeout
    # budget: confirmed empirically that `evaluate=False` on
    # `parse_expr` alone does NOT stop a nested eager-evaluating call
    # like `factorial(factorial(20))` from hanging during `.subs()`/
    # `bool()` -- SymPy's automatic simplification of concrete-integer
    # subexpressions can still fire on those later steps regardless of
    # the flag passed at parse time.
    try:
        with _bounded_parse_time():
            try:
                expr = parse_expr(
                    expression,
                    local_dict=local_dict,
                    global_dict=_SAFE_GLOBAL_DICT,
                    transformations=_SYMPY_TRANSFORMATIONS,
                    evaluate=False,
                )
            except (SyntaxError, TypeError, sympy.SympifyError) as exc:
                raise EquationParseError(
                    f"Could not parse constraint '{expression}': {exc}"
                ) from exc

            substitutions = {sympy.Symbol(name): value for name, value in values.items()}
            result = expr.subs(substitutions)
            if isinstance(
                result, (sympy.logic.boolalg.BooleanTrue, sympy.logic.boolalg.BooleanFalse)
            ):
                return bool(result)
            try:
                return bool(result)
            except TypeError as exc:
                raise EquationParseError(
                    f"Constraint '{expression}' could not be evaluated with the given values "
                    "(missing variable?)."
                ) from exc
    except RecursionError as exc:
        raise EquationParseError(
            f"Constraint '{expression}' is too complex to evaluate."
        ) from exc


# ---------------------------------------------------------------------------
# Test case import (Part G) -- JSON / CSV / XLSX, structured rows only.
# Deliberately NOT routed through the document extraction pipeline
# (extract -> chunk -> embed): these are structured numeric imports, not
# prose to search, and reusing that pipeline would silently prose-ify
# numeric test-case rows.
# ---------------------------------------------------------------------------


def _row_to_test_case(row_index: int, row: dict[str, str]) -> ExperimentTestCase:
    inputs: dict[str, str] = {}
    expected_outputs: list[ExperimentOutputValue] = []
    for key, value in row.items():
        if key is None:
            continue
        key = key.strip()
        if value is None:
            continue
        value = str(value).strip()
        if key.lower().startswith("expected_"):
            expected_outputs.append(
                ExperimentOutputValue(
                    output_name=key[len("expected_"):], value=value, kind=OutputKind.EXPECTED
                )
            )
        elif key:
            inputs[key] = value
    return ExperimentTestCase(
        id=f"case-{row_index}", inputs=inputs, expected_outputs=expected_outputs
    )


def parse_test_cases_json(content: bytes) -> tuple[list[ExperimentTestCase], list[dict[str, str]]]:
    imported: list[ExperimentTestCase] = []
    rejected: list[dict[str, str]] = []
    try:
        data = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        return [], [{"row": "?", "reason": f"Invalid JSON: {exc}"}]

    rows = data if isinstance(data, list) else [data]
    for i, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            rejected.append({"row": str(i), "reason": "Row is not a JSON object."})
            continue
        try:
            imported.append(_row_to_test_case(i, {k: str(v) for k, v in row.items()}))
        except (ValueError, TypeError) as exc:
            rejected.append({"row": str(i), "reason": str(exc)})
    return imported, rejected


def parse_test_cases_csv(content: bytes) -> tuple[list[ExperimentTestCase], list[dict[str, str]]]:
    imported: list[ExperimentTestCase] = []
    rejected: list[dict[str, str]] = []
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        return [], [{"row": "?", "reason": f"Not valid UTF-8: {exc}"}]

    reader = csv.DictReader(io.StringIO(decoded))
    for i, row in enumerate(reader, start=1):
        if not any((v or "").strip() for v in row.values()):
            rejected.append({"row": str(i), "reason": "Empty row."})
            continue
        try:
            imported.append(_row_to_test_case(i, row))
        except (ValueError, TypeError) as exc:
            rejected.append({"row": str(i), "reason": str(exc)})
    return imported, rejected


def parse_test_cases_xlsx(content: bytes) -> tuple[list[ExperimentTestCase], list[dict[str, str]]]:
    imported: list[ExperimentTestCase] = []
    rejected: list[dict[str, str]] = []
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - any openpyxl failure is a parse rejection
        return [], [{"row": "?", "reason": f"Could not parse XLSX: {exc}"}]

    sheet = workbook.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        headers = [str(c) if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return [], [{"row": "?", "reason": "Sheet is empty."}]

    for i, row in enumerate(rows_iter, start=1):
        if all(cell is None for cell in row):
            rejected.append({"row": str(i), "reason": "Empty row."})
            continue
        row_dict = {h: ("" if v is None else str(v)) for h, v in zip(headers, row) if h}
        try:
            imported.append(_row_to_test_case(i, row_dict))
        except (ValueError, TypeError) as exc:
            rejected.append({"row": str(i), "reason": str(exc)})
    return imported, rejected


_MIME_PARSERS = {
    "application/json": parse_test_cases_json,
    "text/csv": parse_test_cases_csv,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": parse_test_cases_xlsx,
}


def parse_test_cases(content: bytes, mime_type: str) -> tuple[list[ExperimentTestCase], list[dict[str, str]]]:
    parser = _MIME_PARSERS.get(mime_type)
    if parser is None:
        raise ValueError(f"Unsupported test-case import format: {mime_type}")
    return parser(content)


# ---------------------------------------------------------------------------
# Visualization data prep (Part H/R). Backend produces raw trace data only
# -- no chart is rendered server-side; the frontend renders it with
# Plotly.js. No causal language is generated here (Part H).
# ---------------------------------------------------------------------------


def prepare_input_output_series(
    test_cases: list[ExperimentTestCase], input_name: str, output_name: str
) -> list[tuple[str, float | None]]:
    """One (x, y) pair per test case that has both fields, x sorted
    numerically when possible. Returns raw pairs; the caller wraps them
    into a `VisualizationSeries` -- kept separate so it stays trivially
    unit-testable without constructing the full response schema."""
    pairs: list[tuple[str, float | None]] = []
    for case in test_cases:
        if input_name not in case.inputs:
            continue
        x_value = case.inputs[input_name]
        y_value: float | None = None
        for out in case.expected_outputs:
            if out.output_name == output_name:
                try:
                    y_value = float(out.value)
                except ValueError:
                    y_value = None
                break
        pairs.append((x_value, y_value))

    def _sort_key(pair: tuple[str, float | None]) -> float:
        try:
            return float(pair[0])
        except ValueError:
            return float("inf")

    return sorted(pairs, key=_sort_key)


def new_plan_id() -> uuid.UUID:
    return uuid.uuid4()


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


# Hard cap on planned combinations (Part J): a sweep is only ever
# DISPLAYED for review, never executed here, but an unbounded Cartesian
# product would still be a real cost (and a real footgun) for whoever
# reviews it -- so this is a display-time safety limit, not an
# execution limit.
MAX_SWEEP_COMBINATIONS = 200


class SweepTooLargeError(ValueError):
    """Raised when a requested sweep would exceed `MAX_SWEEP_COMBINATIONS`."""


def plan_parameter_sweep(
    parameter_values: dict[str, list[str]], *, existing_variant_ids: set[str] | None = None
) -> list[ExperimentVariant]:
    """Expand a sweep spec (Part J) into PLANNED variants -- never
    executed. Raises `SweepTooLargeError` rather than silently
    truncating the combination set, since a silently-truncated sweep
    would misrepresent what the user asked to review.
    """
    if not parameter_values:
        return []
    names = list(parameter_values.keys())
    value_lists = [parameter_values[name] for name in names]
    total = 1
    for values in value_lists:
        total *= max(len(values), 1)
    if total > MAX_SWEEP_COMBINATIONS:
        raise SweepTooLargeError(
            f"This sweep would plan {total} combinations, exceeding the review limit of "
            f"{MAX_SWEEP_COMBINATIONS}. Narrow the parameter ranges before planning."
        )

    existing_variant_ids = existing_variant_ids or set()
    variants: list[ExperimentVariant] = []
    for combo in itertools.product(*value_lists):
        overrides = dict(zip(names, combo))
        label = ", ".join(f"{k}={v}" for k, v in overrides.items())
        variant_id = f"sweep-{uuid.uuid4().hex[:8]}"
        while variant_id in existing_variant_ids:
            variant_id = f"sweep-{uuid.uuid4().hex[:8]}"
        existing_variant_ids.add(variant_id)
        variants.append(
            ExperimentVariant(
                id=variant_id,
                name=label,
                is_baseline=False,
                overrides=overrides,
                status=RunPlanStatus.PLANNED,
            )
        )
    return variants
